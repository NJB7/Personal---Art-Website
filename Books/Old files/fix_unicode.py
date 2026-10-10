"""Fixes garbled accents ("mojibake") in a book export, e.g.  PÃ©rez -> Pérez,  CamÃ©lias -> Camélias.

This happens when text saved as UTF-8 gets read as Windows-1252/Latin-1 somewhere along the way.

Usage (in a terminal, beside the file):
    python fix_unicode.py                                  # uses the default file name below
    python fix_unicode.py "my_export.xlsx"                 # or name the file
    python fix_unicode.py "goodreads_library_export.csv"   # .csv files work too

It never changes your original file. It writes a copy ending in _fixed and prints every change.
Install once if needed:  pip install openpyxl
"""
import csv
import os
import re
import sys

DEFAULT_FILE = "wordpress_book_database_import_review.xlsx"

# Characters that start a garbled sequence. Text without any of these is left alone.
SUSPECT = re.compile("[ÂÃÅâƒ‘-„†-•…‰‹›€™]")


def to_bytes(text):
    """Turn each character back into the single byte it was mistaken for (Windows-1252, else Latin-1)."""
    out = bytearray()
    for ch in text:
        try:
            out += ch.encode("cp1252")
        except UnicodeEncodeError:
            if ord(ch) < 256:
                out.append(ord(ch))      # the few bytes Windows-1252 leaves undefined
            else:
                raise
    return bytes(out)


def fix_text(text):
    """Return the repaired text, or the original if it does not look garbled / cannot be repaired safely."""
    current = text
    for _ in range(3):                   # a few files were garbled twice
        if not SUSPECT.search(current):
            break
        try:
            repaired = to_bytes(current).decode("utf-8")
        except (UnicodeEncodeError, UnicodeDecodeError):
            break
        if repaired == current:
            break
        current = repaired
    return current


def report(changes):
    print("\n%d value(s) fixed." % len(changes))
    for where, before, after in changes[:200]:
        print("  %-28s %r -> %r" % (where, before[:60], after[:60]))
    if len(changes) > 200:
        print("  ... and %d more" % (len(changes) - 200))


def fix_xlsx(path, out_path):
    from openpyxl import load_workbook
    wb = load_workbook(path)             # keeps formatting and formulas
    changes = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and not str(cell.value).startswith("="):
                    fixed = fix_text(cell.value)
                    if fixed != cell.value:
                        changes.append(("%s!%s" % (ws.title, cell.coordinate), cell.value, fixed))
                        cell.value = fixed
    wb.save(out_path)
    return changes


def fix_csv(path, out_path):
    changes = []
    with open(path, newline="", encoding="utf-8-sig") as src:
        rows = list(csv.reader(src))
    for r, row in enumerate(rows, 1):
        for c, value in enumerate(row):
            fixed = fix_text(value)
            if fixed != value:
                changes.append(("row %d, column %d" % (r, c + 1), value, fixed))
                row[c] = fixed
    with open(out_path, "w", newline="", encoding="utf-8-sig") as dst:   # utf-8-sig so Excel reads accents correctly
        csv.writer(dst).writerows(rows)
    return changes


def main():
    path = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_FILE
    if not os.path.exists(path):
        sys.exit("Could not find %r. Put this script in the same folder as the file, or give the file name." % path)
    base, ext = os.path.splitext(path)
    out_path = base + "_fixed" + ext
    if ext.lower() == ".csv":
        changes = fix_csv(path, out_path)
    elif ext.lower() in (".xlsx", ".xlsm"):
        changes = fix_xlsx(path, out_path)
    else:
        sys.exit("Only .xlsx and .csv files are supported.")
    report(changes)
    print("\nSaved %s" % out_path)


if __name__ == "__main__":
    main()
