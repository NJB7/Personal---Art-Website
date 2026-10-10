# Book database enrichment - plain-Python version of enrich_book_database.ipynb
# Install once:  pip install requests openpyxl pillow
# The '# %%' lines mark cells: VS Code, Spyder and PyCharm run them as notebook cells,
# or just run the whole file with:  python enrich_book_database.py
# Edit the settings block near the top first (CONTACT_EMAIL is required).

# %% [markdown]
# # Book database: clean up and enrich the Goodreads import
#
# This notebook reads your Goodreads workbook, tidies it, looks up what is missing, downloads cover images,
# and writes **`wordpress_book_database_import_review.xlsx`** for you to check before importing.
#
# **What it does**
# - Normalises Format to `Print` / `Audiobook` / `eBook`, Gender to `Male` / `Female` / `Other`, Fiction/Nonfiction, ratings, flags and country spellings.
# - Moves each Review from the Books sheet onto the matching Read Instance.
# - Fills blanks only. **Your existing values are never overwritten.**
# - Looks up pages, publication year, Fiction/Nonfiction and original language for books (Open Library, and optionally Google Books).
# - Looks up author country and gender (Open Library to Wikidata).
# - Downloads small JPEG covers from Open Library, one per book.
# - Leaves **Genre, BIPOC, LGBTQ+ and Disabled/Neurodiverse** alone.
#
# **Colours in the review workbook**
# - **Yellow**: looked up on the web.
# - **Orange**: a guess or inference, so please check it.
# - **Blue**: moved or derived from your own data.
# - The `Review Notes` column on each row says where a value came from. The `Problems` sheet lists things that need a decision.
#
# **Before you start:** save this notebook in the same folder as your Goodreads workbook, and install Jupyter if you haven't.
#
# **Running it:** fill in the settings cell, then *Run All*. A full run typically takes an hour or so, mostly waiting on
# web services (the pace is deliberately polite). Everything fetched is cached in a small database file next to the
# notebook, so you can stop and re-run at any time without repeating work. Set `TEST_LIMIT = 40` first for a quick trial.

# %% [markdown]
# ## 1. Settings

# %%
INPUT_XLSX    = "wordpress_book_database_import_goodreads_isbn_language_country.xlsx"
OUTPUT_XLSX   = "wordpress_book_database_import_review.xlsx"

# Open Library asks API users to identify themselves (this gives you 3 requests/second instead of 1).
CONTACT_EMAIL = "nicole.jayne.bergen@gmail.com"          # <-- put your real email here

# Optional. A free key from the Google Cloud console (Books API). Used only to fill Fiction/Nonfiction
# for books Open Library could not classify. Leave as None to skip Google entirely.
GOOGLE_BOOKS_API_KEY      = None
GOOGLE_BOOKS_MAX_PER_RUN  = 900            # the free daily quota is about 1,000; re-run tomorrow for the rest

DOWNLOAD_COVERS    = True
COVER_DIR          = "covers"              # files are saved as covers/<ImportBookID>.jpg
COVER_MAX_WIDTH    = 400                   # pixels; covers are shrunk to this width
COVER_JPEG_QUALITY = 80
COVER_MAX_KB       = 60                    # each file is squeezed under this size
MAKE_COVERS_ZIP    = True                  # also writes covers.zip for easy upload

INFER_READ_FORMAT_FROM_BOOK = True         # blank Read Instance format <- the book's format (orange)
INFER_LANGUAGE_READ         = True         # blank Language Read <- "English" when the book was originally English (orange)

TEST_LIMIT = None                          # e.g. 40 to try only the first 40 books
CACHE_DB   = "enrichment_cache.sqlite"

# %% [markdown]
# ## 2. Helpers and rules

# %%
import os, re, io, json, time, sqlite3, zipfile, hashlib, datetime, threading, unicodedata, collections
from difflib import SequenceMatcher
from concurrent.futures import ThreadPoolExecutor

import requests
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

if "example.com" in CONTACT_EMAIL:
    raise RuntimeError("Please set CONTACT_EMAIL in the settings cell to your real email address.")

RUN_STARTED = datetime.datetime.now()
THIS_YEAR = RUN_STARTED.year

# ---------- allowed values (must match the WordPress snippets) ----------
ALLOWED_FORMATS = {"Print", "Audiobook", "eBook"}
ALLOWED_GENDERS = {"Male", "Female", "Other"}
ALLOWED_FNF     = {"Fiction", "Nonfiction"}
ALLOWED_RATINGS = {"1", "1.5", "2", "2.5", "3", "3.5", "4", "4.5", "5"}

# ---------- format rules: first matching keyword wins ----------
FORMAT_RULES = [
    ("Audiobook", ("audio",)),
    ("eBook",     ("kindle", "ebook", "e-book", "epub", "nook", "digital")),
    ("Print",     ("paper", "hardcover", "hardback", "binding", "bound", "pocket", "board",
                   "bunko", "leather", "diary", "unbound", "print", "mass market")),
]

# ---------- country spellings ----------
COUNTRY_ALIASES = {
    "us": "US", "usa": "US", "u.s.": "US", "u.s.a.": "US", "united states": "US",
    "united states of america": "US", "america": "US",
    "uk": "UK", "u.k.": "UK", "united kingdom": "UK", "great britain": "UK",
    "kingdom of great britain": "UK", "united kingdom of great britain and ireland": "UK",
    # typos found in the spreadsheet
    "vietname": "Vietnam", "palastine": "Palestine", "guatamala": "Guatemala", "haita": "Haiti", "ca": "Canada",
    # Wikidata labels -> the names you already use
    "czechia": "Czech Republic", "state of palestine": "Palestine", "people's republic of china": "China",
    "republic of china": "Taiwan", "united arab emirates": "UAE", "ivory coast": "Côte d'Ivoire",
    "türkiye": "Turkey", "turkiye": "Turkey", "east timor": "Timor-Leste",
}
ENGLISH_SPEAKING = {"US", "UK", "Canada", "Australia", "New Zealand", "Ireland", "Scotland", "England", "Wales"}

# ---------- language names ----------
LANGUAGE_ALIASES = {"f": "French", "farsi": "Persian"}
MARC_LANGUAGES = {
    "eng": "English", "fre": "French", "fra": "French", "spa": "Spanish", "ger": "German", "deu": "German",
    "ita": "Italian", "por": "Portuguese", "rus": "Russian", "jpn": "Japanese", "chi": "Chinese", "zho": "Chinese",
    "ara": "Arabic", "heb": "Hebrew", "hin": "Hindi", "kor": "Korean", "dut": "Dutch", "nld": "Dutch",
    "swe": "Swedish", "nor": "Norwegian", "dan": "Danish", "fin": "Finnish", "ice": "Icelandic", "isl": "Icelandic",
    "pol": "Polish", "cze": "Czech", "ces": "Czech", "hun": "Hungarian", "gre": "Greek", "ell": "Greek",
    "grc": "Ancient Greek", "lat": "Latin", "tur": "Turkish", "per": "Persian", "fas": "Persian", "urd": "Urdu",
    "ben": "Bengali", "tam": "Tamil", "tha": "Thai", "vie": "Vietnamese", "ind": "Indonesian", "rum": "Romanian",
    "ron": "Romanian", "ukr": "Ukrainian", "bul": "Bulgarian", "srp": "Serbian", "scc": "Serbian", "hrv": "Croatian",
    "scr": "Croatian", "cat": "Catalan", "gle": "Irish", "wel": "Welsh", "cym": "Welsh", "afr": "Afrikaans",
    "yid": "Yiddish", "swa": "Swahili", "est": "Estonian", "lit": "Lithuanian", "lav": "Latvian", "slv": "Slovenian",
    "slo": "Slovak", "slk": "Slovak", "uzb": "Uzbek", "kaz": "Kazakh", "mon": "Mongolian", "nob": "Norwegian",
    "nno": "Norwegian", "enm": "Middle English", "ang": "Old English", "non": "Old Norse", "cop": "Coptic",
    "san": "Sanskrit", "tib": "Tibetan", "bod": "Tibetan", "gla": "Scottish Gaelic", "eus": "Basque", "baq": "Basque",
}

# ---------- Fiction / Nonfiction signals (Open Library subjects and Google categories) ----------
FICTION_EXACT = {"fiction", "novels", "short stories", "juvenile fiction", "fantasy fiction", "science fiction",
                 "historical fiction", "mystery fiction", "detective and mystery stories", "romance fiction",
                 "horror fiction", "adventure stories", "young adult fiction", "literary fiction", "dystopian fiction",
                 "thrillers", "fairy tales", "folklore", "picture books"}
NONFICTION_EXACT = {"nonfiction", "non-fiction", "biography", "autobiography", "memoir", "memoirs", "history",
                    "self-help", "psychology", "philosophy", "religion", "science", "social science", "true crime",
                    "cooking", "travel", "politics and government", "essays", "reference", "biographies",
                    "spirituality", "occultism", "witchcraft", "magic", "mythology", "nature", "history and criticism",
                    "literary criticism", "education", "business", "economics", "sociology", "health"}
GOOGLE_NONFICTION = ("biography", "autobiography", "history", "religion", "self-help", "science", "psychology",
                     "philosophy", "social science", "true crime", "cooking", "travel", "political science",
                     "nature", "body, mind", "health", "business", "education", "reference", "literary criticism",
                     "juvenile nonfiction", "young adult nonfiction", "family & relationships", "language arts",
                     "technology", "medical", "computers", "mathematics", "humor/essays")

WRITER_WORDS = re.compile(r"writer|author|novelist|poet|playwright|dramatist|journalist|essayist|historian|biographer|"
                          r"critic|screenwriter|memoirist|columnist|translator|philosopher|scholar|academic|professor|"
                          r"scientist|theologian|illustrator|cartoonist|lyricist|activist|blogger|comics", re.I)

# ---------- output colours ----------
FILL_WEB     = PatternFill("solid", fgColor="FFF2CC")   # yellow
FILL_GUESS   = PatternFill("solid", fgColor="F8CBAD")   # orange
FILL_DERIVED = PatternFill("solid", fgColor="DDEBF7")   # blue
KIND_FILL = {"web": FILL_WEB, "guess": FILL_GUESS, "derived": FILL_DERIVED}
KIND_LABEL = {"web": "looked up (yellow)", "guess": "guess - check (orange)", "derived": "from your data (blue)"}


def blank(v):
    return v is None or (isinstance(v, str) and v.strip() == "")


def clean_text(v):
    if blank(v):
        return None
    return re.sub(r"[ \t\xa0]+", " ", str(v)).strip()


def clean_line(v):
    """Titles and names: collapse every kind of whitespace (including line breaks) to single spaces."""
    if blank(v):
        return None
    return re.sub(r"\s+", " ", str(v)).strip()


def to_int(v):
    if blank(v):
        return None
    try:
        return int(float(str(v).strip()))
    except ValueError:
        return None


def strip_accents(s):
    return "".join(c for c in unicodedata.normalize("NFKD", s) if not unicodedata.combining(c))


def norm_name(s):
    s = strip_accents(str(s)).lower()
    s = re.sub(r"[^a-z0-9]+", " ", s)
    return re.sub(r"\s+", " ", s).strip()


def names_match(a, b):
    a, b = norm_name(a), norm_name(b)
    if not a or not b:
        return False
    if a == b or a.replace(" ", "") == b.replace(" ", ""):
        return True
    return SequenceMatcher(None, a.replace(" ", ""), b.replace(" ", "")).ratio() >= 0.92


def clean_title(t):
    t = str(t).replace("\n", " ")
    t = re.sub(r"\([^)]*\)", " ", t)          # series text in brackets
    t = t.split(":")[0]                          # drop subtitle
    return re.sub(r"\s+", " ", t).strip()


def title_similarity(a, b):
    return SequenceMatcher(None, norm_name(clean_title(a)), norm_name(clean_title(b))).ratio()


# ---------- ISBN ----------
def clean_isbn(v):
    if blank(v):
        return None
    s = re.sub(r"[^0-9Xx]", "", str(v)).upper()
    return s or None


def isbn10_to_13(isbn10):
    core = "978" + isbn10[:9]
    total = sum(int(d) * (1 if i % 2 == 0 else 3) for i, d in enumerate(core))
    return core + str((10 - total % 10) % 10)


def isbn13_valid(s):
    if not s or len(s) != 13 or not s.isdigit():
        return False
    return sum(int(d) * (1 if i % 2 == 0 else 3) for i, d in enumerate(s)) % 10 == 0


def to_isbn13(v):
    s = clean_isbn(v)
    if not s:
        return None
    if len(s) == 13 and isbn13_valid(s):
        return s
    if len(s) == 10:
        return isbn10_to_13(s)
    return None


# ---------- value normalisers ----------
def norm_format(v):
    """Returns (canonical value or None, unrecognised?)"""
    if blank(v):
        return None, False
    s = str(v).strip().lower()
    for canon, words in FORMAT_RULES:
        if any(w in s for w in words):
            return canon, False
    return None, True


def norm_country(v):
    """Returns (value, odd?) - odd means it has brackets or slashes you should look at."""
    if blank(v):
        return None, False
    out, odd = [], False
    for tok in str(v).split(","):
        tok = tok.strip()
        if not tok:
            continue
        if "(" in tok or "/" in tok:
            odd = True
        tok = COUNTRY_ALIASES.get(tok.lower(), tok)
        if tok not in out:
            out.append(tok)
    return (", ".join(out) or None), odd


def norm_gender(v):
    if blank(v):
        return None, False
    m = {"f": "Female", "female": "Female", "m": "Male", "male": "Male", "o": "Other", "other": "Other",
         "nb": "Other", "non-binary": "Other", "nonbinary": "Other"}
    r = m.get(str(v).strip().lower())
    return r, r is None


def norm_fnf(v):
    if blank(v):
        return None, False
    m = {"f": "Fiction", "fiction": "Fiction", "n": "Nonfiction", "nonfiction": "Nonfiction", "non-fiction": "Nonfiction"}
    r = m.get(str(v).strip().lower())
    return r, r is None


def norm_flag(v):
    if blank(v):
        return None
    s = str(v).strip().lower()
    if s in ("1", "1.0", "yes", "y", "true"):
        return "1"
    if s in ("0", "0.0", "no", "n", "false"):
        return "0"
    return str(v).strip()


def norm_rating(v):
    """Returns (rating string or None, was_zero?)"""
    if blank(v):
        return None, False
    try:
        x = float(str(v).strip())
    except ValueError:
        return None, False
    if x == 0:
        return None, True
    x = round(x * 2) / 2
    x = min(max(x, 1), 5)
    return ("%g" % x), False


def norm_date(v):
    """Returns (YYYY-MM-DD or None, invalid?)"""
    if blank(v):
        return None, False
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.strftime("%Y-%m-%d"), False
    s = str(v).strip()
    try:
        datetime.datetime.strptime(s, "%Y-%m-%d")
        return s, False
    except ValueError:
        return None, True


def norm_language_list(v):
    if blank(v):
        return None
    toks = []
    for t in str(v).split(","):
        t = t.strip()
        if t:
            t = LANGUAGE_ALIASES.get(t.lower(), t)
            if t not in toks:
                toks.append(t)
    return ", ".join(toks) or None


def marc_to_name(code):
    return MARC_LANGUAGES.get(str(code).lower())


# %% [markdown]
# ## 3. Network layer (cached, rate limited, polite)

# %%
class QuotaExceeded(Exception):
    pass


class Net:
    MIN_GAP = {"openlibrary.org": 0.4, "covers.openlibrary.org": 0.05, "query.wikidata.org": 1.0, "www.googleapis.com": 0.25}

    def __init__(self, cache_db, email):
        self.db = sqlite3.connect(cache_db, check_same_thread=False)
        self.db.execute("CREATE TABLE IF NOT EXISTS cache (key TEXT PRIMARY KEY, value TEXT)")
        self.db.commit()
        self.lock = threading.Lock()
        self.last = {}
        self.session = requests.Session()
        self.session.headers["User-Agent"] = "BookDatabaseEnricher/1.0 (%s)" % email
        self.stats = collections.Counter()

    # --- cache ---
    def _cache_get(self, key):
        with self.lock:
            row = self.db.execute("SELECT value FROM cache WHERE key=?", (key,)).fetchone()
        return json.loads(row[0]) if row else None

    def _cache_put(self, key, value):
        with self.lock:
            self.db.execute("INSERT OR REPLACE INTO cache (key, value) VALUES (?, ?)", (key, json.dumps(value)))
            self.db.commit()

    def _throttle(self, url):
        host = re.sub(r"^https?://", "", url).split("/")[0]
        gap = self.MIN_GAP.get(host, 0.3)
        with self.lock:
            wait = self.last.get(host, 0) + gap - time.time()
            self.last[host] = max(time.time(), self.last.get(host, 0) + gap)
        if wait > 0:
            time.sleep(wait)

    def _request(self, url, params=None, headers=None, tries=4):
        last_error = None
        for attempt in range(tries):
            self._throttle(url)
            try:
                r = self.session.get(url, params=params, headers=headers, timeout=40)
            except (requests.ConnectionError, requests.Timeout) as e:
                last_error = e
                time.sleep(2 ** attempt)
                continue
            self.stats["requests"] += 1
            if "googleapis.com" in url and r.status_code in (403, 429):
                raise QuotaExceeded("Google Books quota reached")
            if r.status_code == 429 or r.status_code >= 500:
                try:
                    wait = float(r.headers.get("Retry-After", 0))
                except ValueError:
                    wait = 0
                time.sleep(max(wait, 2 ** attempt * 2))
                last_error = RuntimeError("HTTP %s" % r.status_code)
                continue
            return r
        raise RuntimeError("Giving up on %s (%s)" % (url, last_error))

    # --- public ---
    def get_json(self, url, params=None, headers=None):
        key = "GET " + url + " " + json.dumps(params or {}, sort_keys=True)
        hit = self._cache_get(key)
        if hit is not None:
            self.stats["cache_hits"] += 1
            return hit
        try:
            r = self._request(url, params, headers)
        except RuntimeError:
            self.stats["failed"] += 1
            return {"__status": -1}                 # not cached, so a re-run retries it
        if r.status_code == 404:
            data = {"__status": 404}
        elif r.status_code == 200:
            try:
                data = r.json()
            except ValueError:
                return {"__status": 0}
        else:
            return {"__status": r.status_code}      # not cached
        self._cache_put(key, data)
        return data

    def get_bytes(self, url):
        try:
            r = self._request(url)
        except RuntimeError:
            return 0, b""
        return r.status_code, r.content


def not_found(data):
    return not isinstance(data, dict) or "__status" in data


def progress(i, n, label):
    step = max(1, -(-n // 10))
    if i == 1 or i == n or i % step == 0:
        print("  %s: %d / %d" % (label, i, n), flush=True)


net = Net(CACHE_DB, CONTACT_EMAIL)
print("Ready. Cache file:", os.path.abspath(CACHE_DB))

# %% [markdown]
# ## 4. Load your workbook and clean it
# Nothing is looked up yet. This step only tidies your own data and records every change.

# %%
def read_sheet(wb, name):
    rows = list(wb[name].iter_rows(values_only=True))
    header = list(rows[0])
    return header, [dict(zip(header, r)) for r in rows[1:] if any(not blank(x) for x in r)]


wb_in = load_workbook(INPUT_XLSX, data_only=True)
AUTH_COLS, authors_raw = read_sheet(wb_in, "Authors")
BOOK_COLS, books_raw = read_sheet(wb_in, "Books")
READ_COLS, reads_raw = read_sheet(wb_in, "Read Instances")

# provenance: (sheet, key, column) -> (kind, note)
prov = {}
row_notes = collections.defaultdict(list)       # (sheet, key) -> [text]
problems = []                                   # dicts: sheet, key, name, issue, details
changes = collections.Counter()                 # tidy-ups applied to your own values
fill_counts = collections.Counter()             # (sheet, column, kind) -> n


def mark(sheet, key, col, kind, note):
    prov[(sheet, key, col)] = (kind, note)
    fill_counts[(sheet, col, kind)] += 1
    if note:
        row_notes[(sheet, key)].append("%s: %s" % (col, note))


def problem(sheet, key, name, issue, details=""):
    problems.append({"Sheet": sheet, "Row ID": key, "Name / Title": name, "Issue": issue, "Details": details})


# ---------- authors ----------
def tidy_line(v, what):
    out = clean_line(v)
    if out is not None and out != v:
        changes["Extra spaces / line breaks removed from %s" % what] += 1
    return out


authors = {}
for r in authors_raw:
    aid = str(r["ImportAuthorID"]).strip()
    country, odd = norm_country(r.get("Country"))
    if country != clean_text(r.get("Country")):
        changes["Author country spelling tidied"] += 1
    if odd:
        problem("Authors", aid, r["Author"], "Country has brackets or a slash", str(r.get("Country")))
    gender, bad = norm_gender(r.get("Gender"))
    if bad:
        problem("Authors", aid, r["Author"], "Gender value not recognised (WordPress allows Male / Female / Other)",
                "Original value: %r - left blank" % r.get("Gender"))
    elif gender and gender != r.get("Gender"):
        changes["Author gender expanded (F/M/O to Female/Male/Other)"] += 1
    authors[aid] = {
        "id": aid, "name": tidy_line(r["Author"], "author names"), "country": country, "country_details": r.get("CountryDetails"),
        "gender": gender, "bipoc": norm_flag(r.get("BIPOC")), "lgbtq": norm_flag(r.get("LGBTQ+")),
        "disabled": norm_flag(r.get("Disabled/Neurodiverse")),
    }

# ---------- books ----------
books = []
for r in books_raw:
    bid = str(r["ImportBookID"]).strip()
    fmt, bad = norm_format(r.get("Format"))
    if bad:
        problem("Books", bid, r["BookTitle"], "Format not recognised", "Original value: %r - left blank" % r.get("Format"))
    elif fmt and fmt != r.get("Format"):
        changes["Book format normalised to Print/Audiobook/eBook"] += 1
    fnf, bad = norm_fnf(r.get("Fiction/Nonfiction"))
    if bad:
        problem("Books", bid, r["BookTitle"], "Fiction/Nonfiction not recognised", repr(r.get("Fiction/Nonfiction")))
    elif fnf and fnf != r.get("Fiction/Nonfiction"):
        changes["Fiction/Nonfiction expanded (F/N)"] += 1
    lang = norm_language_list(r.get("OriginalLanguage"))
    if lang != clean_text(r.get("OriginalLanguage")):
        changes["Original language name tidied (e.g. Farsi to Persian)"] += 1
    isbn_clean = clean_isbn(r.get("ISBN"))
    books.append({
        "id": bid, "title": tidy_line(r.get("BookTitle"), "book titles"), "author_name": clean_line(r.get("Author")),
        "author_id": str(r.get("ImportAuthorID")).strip() if not blank(r.get("ImportAuthorID")) else None,
        "pages": to_int(r.get("Pages")), "format": fmt, "fnf": fnf, "genre": r.get("Genre"),
        "orig_lang": lang, "isbn": isbn_clean, "isbn13": to_isbn13(isbn_clean), "year": to_int(r.get("PublicationYear")),
        "review": clean_text(r.get("Review")), "country_col": r.get("Country"), "status": None,
    })
    if isbn_clean and not to_isbn13(isbn_clean):
        problem("Books", bid, r["BookTitle"], "ISBN is not a valid ISBN-10/13", str(r.get("ISBN")))

# ---------- read instances ----------
reads = []
for i, r in enumerate(reads_raw):
    key = str(i + 2)                    # the row number in your original sheet
    fmt, bad = norm_format(r.get("Format"))
    if bad:
        problem("Read Instances", key, r["BookTitle"], "Format not recognised", "Original value: %r - left blank" % r.get("Format"))
    elif fmt and fmt != r.get("Format"):
        changes["Read format normalised to Print/Audiobook/eBook"] += 1
    rating, was_zero = norm_rating(r.get("Rating"))
    if was_zero:
        changes["Rating of 0 (unrated) left blank"] += 1
    elif rating is not None and rating != str(r.get("Rating")):
        changes["Rating tidied (4.0 to 4)"] += 1
    date, invalid = norm_date(r.get("DateRead"))
    if invalid:
        problem("Read Instances", key, r["BookTitle"], "DateRead is not a valid YYYY-MM-DD date", repr(r.get("DateRead")))
    lang = norm_language_list(r.get("LanguageRead"))
    if lang != clean_text(r.get("LanguageRead")):
        changes["Language Read tidied (F to French)"] += 1
    reads.append({
        "key": key, "book_id": str(r.get("Book")).strip(), "title": tidy_line(r.get("BookTitle"), "book titles"), "date": date,
        "rating": rating, "format": fmt, "lang": lang, "year": r.get("Year"), "read_flag": r.get("Read"), "review": None,
    })

books_by_id = {b["id"]: b for b in books}
reads_by_book = collections.defaultdict(list)
for rd in reads:
    reads_by_book[rd["book_id"]].append(rd)

changes["Pages and PublicationYear stored as text turned into numbers"] = sum(1 for r in books_raw if isinstance(r.get("Pages"), str) or isinstance(r.get("PublicationYear"), str))
if not changes["Pages and PublicationYear stored as text turned into numbers"]:
    del changes["Pages and PublicationYear stored as text turned into numbers"]

# ---------- reviews: Books sheet -> Read Instances ----------
for b in books:
    if not b["review"]:
        continue
    rs = reads_by_book.get(b["id"], [])
    if not rs:
        problem("Books", b["id"], b["title"], "Has a review but no Read Instance to attach it to", b["review"])
        continue
    target = max(rs, key=lambda x: (x["date"] or "", int(x["key"])))
    target["review"] = b["review"]
    note = "moved from the Books sheet"
    if len(rs) > 1:
        note += " (this book has %d reads; attached to the latest)" % len(rs)
    mark("Read Instances", target["key"], "Review", "derived", note)

# ---------- status: a book with no read instance is probably 'To Read' ----------
for b in books:
    b["status"] = "Read" if b["id"] in reads_by_book else "To Read"
    mark("Books", b["id"], "Status", "derived" if b["status"] == "Read" else "guess", "")

# ---------- duplicate ISBNs ----------
by_isbn = collections.defaultdict(list)
for b in books:
    if b["isbn13"]:
        by_isbn[b["isbn13"]].append(b)
dup_groups = {k: v for k, v in by_isbn.items() if len(v) > 1}
for isbn, group in dup_groups.items():
    for b in group:
        problem("Books", b["id"], b["title"], "Same ISBN as another book record (possible duplicate)",
                "ISBN %s also on: %s" % (isbn, "; ".join("%s (%s)" % (o["title"], o["id"]) for o in group if o is not b)))

if TEST_LIMIT:
    print("*** TEST RUN: web lookups are limited to the first %d books. Set TEST_LIMIT = None for the full run. ***" % TEST_LIMIT)
print("Authors: %d | Books: %d | Read Instances: %d" % (len(authors), len(books), len(reads)))
print("Reviews moved:", sum(1 for r in reads if r["review"]), "| books with duplicate ISBNs:", sum(len(v) for v in dup_groups.values()))

# %% [markdown]
# ## 5. Open Library: look up books by ISBN (batched)

# %%
OL_FIELDS = "key,title,author_key,author_name,first_publish_year,number_of_pages_median,subject,isbn,cover_i"
OL_SEARCH = "https://openlibrary.org/search.json"


def trim_doc(doc):
    return {
        "key": doc.get("key"), "title": doc.get("title"),
        "author_key": doc.get("author_key", []), "author_name": doc.get("author_name", []),
        "first_publish_year": doc.get("first_publish_year"),
        "pages_median": doc.get("number_of_pages_median"),
        "subjects": [s.lower() for s in doc.get("subject", [])[:80]],
        "cover_i": doc.get("cover_i"),
    }


def ol_search_batch(isbns):
    q = "isbn:(" + " OR ".join(isbns) + ")"
    data = net.get_json(OL_SEARCH, {"q": q, "fields": OL_FIELDS, "limit": 100})
    return data.get("docs", []) if isinstance(data, dict) else []


def run_isbn_lookup(target_books, batch_size=40):
    wanted = {}                                          # isbn (13 or original 10) -> [book ids]
    for b in target_books:
        for form in {b["isbn13"], b["isbn"]}:
            if form:
                wanted.setdefault(form, []).append(b["id"])
    isbn13s = sorted({b["isbn13"] for b in target_books if b["isbn13"]})
    found = {}                                           # book id -> trimmed doc
    batches = [isbn13s[i:i + batch_size] for i in range(0, len(isbn13s), batch_size)]
    for n, batch in enumerate(batches, 1):
        progress(n, len(batches), "Open Library batches")
        failed_before = net.stats["failed"]
        docs = ol_search_batch(batch)
        if not docs and len(batch) > 1 and net.stats["failed"] == failed_before:   # OR query not understood? one by one
            for single in batch:
                docs.extend(ol_search_batch([single]))
        for doc in docs:
            for isbn in doc.get("isbn", []):
                for bid in wanted.get(isbn, []):
                    found.setdefault(bid, trim_doc(doc))
    return found


targets = books[:TEST_LIMIT] if TEST_LIMIT else books

# Quick check (a few seconds) so a connection or query problem shows up now, not an hour from now.
_sample = [b for b in books if b["isbn13"]][:3]
_before = net.stats["failed"]
_found = run_isbn_lookup(_sample)
if net.stats["failed"] > _before:
    raise RuntimeError("Could not reach Open Library. Check your internet connection and run this cell again.")
print("Smoke test: Open Library recognised %d of %d sample ISBNs." % (len({b["isbn13"] for b in _sample if b["id"] in _found}), len(_sample)))
if not _found:
    raise RuntimeError("Open Library answered but did not recognise any of your first ISBNs. Try again later; "
                     "if it keeps happening the search format may have changed.")

print("Looking up %d books by ISBN..." % sum(1 for b in targets if b["isbn13"]))
ol_docs = run_isbn_lookup([b for b in targets if b["isbn13"]])
print("Found on Open Library by ISBN: %d of %d" % (len(ol_docs), sum(1 for b in targets if b["isbn13"])))

# %% [markdown]
# ## 6. Open Library: title + author search for books with no ISBN match
# Matches here are less certain, so anything they fill is shown in orange.

# %%
def ol_search_title(title, author):
    data = net.get_json(OL_SEARCH, {"title": clean_title(title), "author": author or "", "fields": OL_FIELDS, "limit": 5})
    best, best_score = None, 0
    for doc in (data.get("docs", []) if isinstance(data, dict) else []):
        score = SequenceMatcher(None, norm_name(clean_title(title)), norm_name(clean_title(doc.get("title", "")))).ratio()
        surname = norm_name(author or "").split(" ")[-1] if author else ""
        names = [norm_name(n) for n in doc.get("author_name", [])]
        if surname and not any(surname in n for n in names):
            continue
        if score > best_score:
            best, best_score = doc, score
    return (trim_doc(best), best_score) if best and best_score >= 0.88 else (None, best_score)


title_matched = {}                      # book id -> score
missing = [b for b in targets if b["id"] not in ol_docs and b["title"]]
print("Trying title/author search for %d books..." % len(missing))
for n, b in enumerate(missing, 1):
    progress(n, len(missing), "title searches")
    doc, score = ol_search_title(b["title"], b["author_name"])
    if doc:
        ol_docs[b["id"]] = doc
        title_matched[b["id"]] = round(score, 2)
print("Matched by title: %d | still unmatched: %d" % (len(title_matched), len(missing) - len(title_matched)))

# %% [markdown]
# ## 7. Open Library: edition details (only for books that still have gaps)
# Gives the edition's own page count, year, and a "translated from" language when OL has one.

# %%
def parse_edition(d):
    if not isinstance(d, dict) or "__status" in d:
        return None
    pages = to_int(d.get("number_of_pages")) or to_int(re.sub(r"\D.*", "", str(d.get("pagination", ""))) or None)
    m = re.search(r"(1[0-9]{3}|20[0-9]{2})", str(d.get("publish_date", "")))
    code = lambda lst: [x["key"].split("/")[-1] for x in lst if isinstance(x, dict) and "key" in x]
    covers = [c for c in d.get("covers", []) if isinstance(c, int) and c > 0]
    return {"pages": pages if isinstance(pages, int) and pages > 0 else None, "year": int(m.group(1)) if m else None,
            "languages": code(d.get("languages", [])), "translated_from": code(d.get("translated_from", [])),
            "cover_id": covers[0] if covers else None}


need_edition = [b for b in targets if (b["isbn13"] or b["isbn"]) and (b["pages"] is None or b["year"] is None or b["orig_lang"] is None)]
editions = {}
print("Looking up %d individual editions..." % len(need_edition))
for n, b in enumerate(need_edition, 1):
    progress(n, len(need_edition), "editions")
    data = net.get_json("https://openlibrary.org/isbn/%s.json" % (b["isbn13"] or b["isbn"]))
    if not_found(data) and b["isbn"] and b["isbn"] != b["isbn13"]:
        data = net.get_json("https://openlibrary.org/isbn/%s.json" % b["isbn"])
    ed = parse_edition(data)
    if ed:
        editions[b["id"]] = ed
print("Editions found: %d" % len(editions))

# %% [markdown]
# ## 8. Google Books (optional): Fiction / Nonfiction for books Open Library could not classify

# %%
def classify_from_subjects(subjects):
    """Returns (value or None, strength 'strong'/'weak')."""
    subs = set(subjects)
    fic = {s for s in subs if s in FICTION_EXACT}
    non = {s for s in subs if s in NONFICTION_EXACT}
    weak_fic = {s for s in subs if s.endswith(" fiction") and "criticism" not in s and "history" not in s}
    if fic and not non:
        return "Fiction", "strong"
    if non and not fic and not weak_fic:
        return "Nonfiction", "strong"
    if weak_fic and not non:
        return "Fiction", "weak"
    return None, None


def classify_from_google(categories):
    cats = [c.lower() for c in categories]
    if any(c.startswith(("fiction", "juvenile fiction", "young adult fiction")) for c in cats):
        return "Fiction"
    if any(any(c.startswith(w) or ("/ " + w) in c for w in GOOGLE_NONFICTION) for c in cats):
        return "Nonfiction"
    return None


def google_lookup(isbn):
    params = {"q": "isbn:%s" % isbn, "maxResults": 1, "fields": "items(volumeInfo(categories))"}
    if GOOGLE_BOOKS_API_KEY:
        params["key"] = GOOGLE_BOOKS_API_KEY
    data = net.get_json("https://www.googleapis.com/books/v1/volumes", params)
    items = data.get("items", []) if isinstance(data, dict) else []
    return items[0].get("volumeInfo", {}).get("categories", []) if items else []


google_fnf = {}
if GOOGLE_BOOKS_API_KEY:
    todo = []
    for b in targets:
        if b["fnf"] is None and b["isbn13"]:
            doc = ol_docs.get(b["id"])
            if not (doc and classify_from_subjects(doc["subjects"])[0]):
                todo.append(b)
    todo.sort(key=lambda b: b["id"] not in reads_by_book)         # books you have read come first
    print("%d books need a Google Books lookup (max %d this run)" % (len(todo), GOOGLE_BOOKS_MAX_PER_RUN))
    done = 0
    try:
        for b in todo:
            before = net.stats["requests"]
            cats = google_lookup(b["isbn13"])
            if net.stats["requests"] > before:
                done += 1
            google_fnf[b["id"]] = classify_from_google(cats)
            progress(len(google_fnf), len(todo), "Google Books")
            if done >= GOOGLE_BOOKS_MAX_PER_RUN:
                print("Reached the per-run limit. Run the notebook again tomorrow to continue; nothing is repeated.")
                break
    except QuotaExceeded:
        print("Google Books daily quota reached. Run the notebook again tomorrow to continue; nothing is repeated.")
else:
    print("No Google Books key set, skipping. (Open Library subjects are still used.)")

# %% [markdown]
# ## 9. Authors: Open Library to Wikidata (country and gender)
# Works from each author's own books: the book's Open Library record points to the author, whose record
# links to Wikidata. Authors that cannot be linked this way are matched by exact name, only when exactly one
# writer has that name, and those matches are shown in orange.

# %%
def sparql(query):
    data = net.get_json("https://query.wikidata.org/sparql", {"query": query, "format": "json"},
                        headers={"Accept": "application/sparql-results+json"})
    return data.get("results", {}).get("bindings", []) if isinstance(data, dict) else []


def sparql_str(s):
    return '"%s"@en' % s.replace("\\", "\\\\").replace('"', '\\"')


def wikidata_facts(qids):
    """qid -> {'gender': set, 'citizenship': set, 'birth_country': set}"""
    facts = {}
    qids = sorted(set(qids))
    batches = [qids[i:i + 80] for i in range(0, len(qids), 80)]
    for n, batch in enumerate(batches, 1):
        progress(n, len(batches), "Wikidata fact batches")
        q = ("SELECT ?a ?genderLabel ?citLabel ?bcLabel WHERE { VALUES ?a { %s } "
             "OPTIONAL { ?a wdt:P21 ?gender . } OPTIONAL { ?a wdt:P27 ?cit . } "
             "OPTIONAL { ?a wdt:P19 ?bp . ?bp wdt:P17 ?bc . } "
             'SERVICE wikibase:label { bd:serviceParam wikibase:language "en" . } }') % " ".join("wd:" + x for x in batch)
        for row in sparql(q):
            qid = row["a"]["value"].rsplit("/", 1)[-1]
            f = facts.setdefault(qid, {"gender": set(), "citizenship": set(), "birth_country": set()})
            for src, dst in (("genderLabel", "gender"), ("citLabel", "citizenship"), ("bcLabel", "birth_country")):
                if src in row and not row[src]["value"].startswith("http") and not re.fullmatch(r"Q\d+", row[src]["value"]):
                    f[dst].add(row[src]["value"])
    return facts


def wikidata_find_by_name(names):
    """name -> set of QIDs of people with a writer-like occupation"""
    out = collections.defaultdict(set)
    names = sorted(set(names))
    batches = [names[i:i + 40] for i in range(0, len(names), 40)]
    for n, batch in enumerate(batches, 1):
        progress(n, len(batches), "Wikidata name batches")
        q = ("SELECT ?name ?a ?occLabel WHERE { VALUES ?name { %s } ?a rdfs:label ?name . ?a wdt:P31 wd:Q5 . "
             "OPTIONAL { ?a wdt:P106 ?occ . } "
             'SERVICE wikibase:label { bd:serviceParam wikibase:language "en" . } }') % " ".join(sparql_str(x) for x in batch)
        for row in sparql(q):
            if WRITER_WORDS.search(row.get("occLabel", {}).get("value", "")):
                out[row["name"]["value"]].add(row["a"]["value"].rsplit("/", 1)[-1])
    return out


# 1) which Open Library author does each of your authors correspond to?
votes = collections.defaultdict(collections.Counter)
for b in targets:
    doc = ol_docs.get(b["id"])
    if not doc or not b["author_id"] or b["author_id"] not in authors:
        continue
    for key, nm in zip(doc["author_key"], doc["author_name"]):
        if names_match(authors[b["author_id"]]["name"], nm):
            votes[b["author_id"]][key] += 1

target_author_ids = {b["author_id"] for b in targets if b["author_id"] in authors}
needs = [a for a in authors.values() if a["id"] in target_author_ids and (a["country"] is None or a["gender"] is None)]
print("%d authors are missing a country or gender" % len(needs))

# 2) Open Library author record -> Wikidata id
qid_for, how = {}, {}
for n, a in enumerate(needs, 1):
    progress(n, len(needs), "Open Library author records")
    if votes[a["id"]]:
        key = votes[a["id"]].most_common(1)[0][0]
        rec = net.get_json("https://openlibrary.org/authors/%s.json" % key)
        qid = (rec.get("remote_ids") or {}).get("wikidata") if isinstance(rec, dict) else None
        if qid and re.fullmatch(r"Q\d+", qid):
            qid_for[a["id"]], how[a["id"]] = qid, "ol-link"

# 3) the rest: exact-name match, only when exactly one writer has that name
rest = [a for a in needs if a["id"] not in qid_for and a["name"] and " & " not in a["name"] and " and " not in a["name"].lower()]
print("Linked via Open Library: %d | trying name match for %d" % (len(qid_for), len(rest)))
found_by_name = wikidata_find_by_name([a["name"] for a in rest])
for a in rest:
    cands = found_by_name.get(a["name"], set())
    if len(cands) == 1:
        qid_for[a["id"]], how[a["id"]] = next(iter(cands)), "name-only"

facts = wikidata_facts(qid_for.values())
print("Wikidata facts retrieved for %d people" % len(facts))


def decide_gender(f):
    mapped = set()
    for g in f["gender"]:
        g = g.lower()
        mapped.add("Male" if g in ("male", "trans man") else "Female" if g in ("female", "trans woman") else "Other")
    return mapped.pop() if len(mapped) == 1 else None


def decide_country(f):
    """Returns (country or None, note, disagreement?)"""
    cit = sorted({COUNTRY_ALIASES.get(c.lower(), c) for c in f["citizenship"]})
    born = sorted({COUNTRY_ALIASES.get(c.lower(), c) for c in f["birth_country"]})
    if len(cit) == 1:
        diff = bool(born) and born != cit
        return cit[0], "citizenship %s%s" % (cit[0], "; born in " + ", ".join(born) if diff else ""), diff
    if len(born) == 1:
        return born[0], "born in %s (citizenship: %s)" % (born[0], ", ".join(cit) or "unknown"), len(cit) > 1
    return None, "", False


author_facts = {}
for aid, qid in qid_for.items():
    f = facts.get(qid)
    if f:
        author_facts[aid] = {"gender": decide_gender(f), "country": decide_country(f), "match": how[aid], "qid": qid}
print("Authors with usable Wikidata results: %d" % len(author_facts))

# %% [markdown]
# ## 10. Download covers from Open Library
# Uses cover ids (which are not rate-limited) rather than ISBN URLs. Images are shrunk to a small JPEG so they
# stay light in WordPress.

# %%
cover_files = {}                      # book id -> relative file path


def save_cover(cover_id, dest):
    from PIL import Image
    status, content = net.get_bytes("https://covers.openlibrary.org/b/id/%s-L.jpg?default=false" % cover_id)
    if status != 200 or len(content) < 2500:
        return False
    try:
        img = Image.open(io.BytesIO(content))
        img.load()
    except Exception:
        return False
    if img.width < 80 or img.height < 80:
        return False
    img = img.convert("RGB")
    resample = Image.Resampling.LANCZOS if hasattr(Image, "Resampling") else Image.LANCZOS
    if img.width > COVER_MAX_WIDTH:
        img = img.resize((COVER_MAX_WIDTH, round(img.height * COVER_MAX_WIDTH / img.width)), resample)
    quality, limit = COVER_JPEG_QUALITY, COVER_MAX_KB * 1024
    while True:
        buf = io.BytesIO()
        img.save(buf, "JPEG", quality=quality, optimize=True, progressive=True)
        if buf.tell() <= limit:
            break
        if quality > 55:
            quality -= 5
        elif img.width > 200:
            img = img.resize((round(img.width * 0.85), round(img.height * 0.85)), resample)
        else:
            break
    with open(dest, "wb") as fh:
        fh.write(buf.getvalue())
    return True


if DOWNLOAD_COVERS:
    os.makedirs(COVER_DIR, exist_ok=True)
    jobs = []
    for b in targets:
        ed, doc = editions.get(b["id"]), ol_docs.get(b["id"])
        cid = (ed or {}).get("cover_id") or (doc or {}).get("cover_i")
        if cid:
            jobs.append((b["id"], cid))
    print("Downloading covers for %d books (already-downloaded files are skipped)..." % len(jobs))

    def do_job(job):
        bid, cid = job
        dest = os.path.join(COVER_DIR, "%s.jpg" % bid)
        if os.path.exists(dest) and os.path.getsize(dest) > 0:
            return bid, dest
        return (bid, dest) if save_cover(cid, dest) else (bid, None)

    with ThreadPoolExecutor(max_workers=4) as pool:
        for n, (bid, dest) in enumerate(pool.map(do_job, jobs), 1):
            progress(n, len(jobs), "covers")
            if dest:
                cover_files[bid] = dest
    total = sum(os.path.getsize(p) for p in cover_files.values())
    print("Covers saved: %d | total %.1f MB | average %.0f KB" % (len(cover_files), total / 1e6, total / max(1, len(cover_files)) / 1e3))
    if MAKE_COVERS_ZIP and cover_files:
        with zipfile.ZipFile("covers.zip", "w", zipfile.ZIP_STORED) as z:
            for p in cover_files.values():
                z.write(p, os.path.basename(p))
        print("Wrote covers.zip")
else:
    print("Cover download switched off.")

# %% [markdown]
# ## 11. Fill the blanks (never overwriting your values)

# %%
# Which year does your PublicationYear column mean: the original year or the edition's year?
agree = [b["year"] == ol_docs[b["id"]]["first_publish_year"] for b in books
         if b["year"] and b["id"] in ol_docs and ol_docs[b["id"]]["first_publish_year"]]
USE_FIRST_PUBLISH_YEAR = (sum(agree) / len(agree) >= 0.5) if len(agree) >= 20 else True
year_rule = ("original (first-publish) year - your existing values match Open Library's first-publish year for %d%% of books" % round(100 * sum(agree) / len(agree))
             if len(agree) >= 20 and USE_FIRST_PUBLISH_YEAR else
             "edition year, falling back to first-publish year" if len(agree) >= 20 else "original (first-publish) year (too few matches to compare)")
print("Publication year rule:", year_rule)


def author_countries(b):
    a = authors.get(b["author_id"])
    if not a:
        return set()
    c = a["country"] or (author_facts.get(a["id"], {}).get("country") or (None,))[0]
    return {x.strip() for x in c.split(",")} if c else set()


# ----- authors -----
for a in authors.values():
    f = author_facts.get(a["id"])
    if not f:
        continue
    kind = "web" if f["match"] == "ol-link" else "guess"
    src = "Wikidata %s%s" % (f["qid"], "" if f["match"] == "ol-link" else " (matched by name only)")
    if a["gender"] is None and f["gender"]:
        a["gender"] = f["gender"]
        mark("Authors", a["id"], "Gender", kind, src)
    country, note, differ = f["country"]
    if a["country"] is None and country:
        a["country"] = country
        mark("Authors", a["id"], "Country", "guess" if (differ or kind == "guess") else "web", "%s: %s" % (src, note))

# ----- books -----
for b in books:
    doc, ed = ol_docs.get(b["id"]), editions.get(b["id"])
    via_title = b["id"] in title_matched
    base_kind = "guess" if via_title else "web"
    src = ("Open Library title match %.0f%%" % (title_matched[b["id"]] * 100)) if via_title else "Open Library"

    if b["pages"] is None:
        if ed and ed["pages"]:
            b["pages"] = ed["pages"]
            mark("Books", b["id"], "Pages", base_kind, src + " edition")
        elif doc and doc["pages_median"]:
            b["pages"] = int(doc["pages_median"])
            mark("Books", b["id"], "Pages", "guess", src + " (median across editions, approximate)")

    if b["year"] is None:
        first = doc["first_publish_year"] if doc else None
        ed_year = ed["year"] if ed else None
        year = (first or ed_year) if USE_FIRST_PUBLISH_YEAR else (ed_year or first)
        if year and 1000 <= year <= THIS_YEAR + 1:
            b["year"] = year
            mark("Books", b["id"], "PublicationYear", base_kind, src)

    if b["fnf"] is None:
        value, strength = (classify_from_subjects(doc["subjects"]) if doc else (None, None))
        if value:
            b["fnf"] = value
            mark("Books", b["id"], "Fiction/Nonfiction", "guess" if (strength == "weak" or via_title) else "web", src + " subjects")
        elif google_fnf.get(b["id"]):
            b["fnf"] = google_fnf[b["id"]]
            mark("Books", b["id"], "Fiction/Nonfiction", "web", "Google Books category")

    if b["orig_lang"] is None:
        eds_langs = [marc_to_name(c) for c in (ed["languages"] if ed else []) if marc_to_name(c)]
        translated = [marc_to_name(c) for c in (ed["translated_from"] if ed else []) if marc_to_name(c)]
        if translated:
            b["orig_lang"] = ", ".join(dict.fromkeys(translated))
            mark("Books", b["id"], "OriginalLanguage", "web", "Open Library edition says translated from this language")
        elif eds_langs and "English" not in eds_langs:
            b["orig_lang"] = eds_langs[0]
            mark("Books", b["id"], "OriginalLanguage", "guess", "edition is in %s with no 'translated from' info; assumed original" % eds_langs[0])
        else:
            countries = author_countries(b)
            if countries and countries <= ENGLISH_SPEAKING:
                b["orig_lang"] = "English"
                mark("Books", b["id"], "OriginalLanguage", "guess", "English assumed (author from %s)" % ", ".join(sorted(countries)))

    if b["id"] in cover_files:
        b["cover"] = "%s/%s.jpg" % (os.path.basename(os.path.normpath(COVER_DIR)), b["id"])
        mark("Books", b["id"], "CoverFile", "guess" if via_title else "web", "")
    else:
        b["cover"] = None

# ----- read instances -----
for rd in reads:
    b = books_by_id.get(rd["book_id"])
    if not b:
        problem("Read Instances", rd["key"], rd["title"], "Book id %s is not on the Books sheet" % rd["book_id"])
        continue
    if rd["format"] is None and INFER_READ_FORMAT_FROM_BOOK and b["format"]:
        rd["format"] = b["format"]
        mark("Read Instances", rd["key"], "Format", "guess", "copied from the book's format")
    if rd["lang"] is None and INFER_LANGUAGE_READ and b["orig_lang"] == "English":
        rd["lang"] = "English"
        mark("Read Instances", rd["key"], "LanguageRead", "guess", "book was originally English")

print("Fills done.", sum(fill_counts.values()), "cells filled or derived.")

# %% [markdown]
# ## 12. Checks, then write the review workbook

# %%
errors = []
for b in books:
    if b["format"] not in (None, *ALLOWED_FORMATS): errors.append(("Books", b["id"], "Format", b["format"]))
    if b["fnf"] not in (None, *ALLOWED_FNF): errors.append(("Books", b["id"], "Fiction/Nonfiction", b["fnf"]))
for rd in reads:
    if rd["format"] not in (None, *ALLOWED_FORMATS): errors.append(("Read Instances", rd["key"], "Format", rd["format"]))
    if rd["rating"] not in (None, *ALLOWED_RATINGS): errors.append(("Read Instances", rd["key"], "Rating", rd["rating"]))
    if rd["date"] and not re.fullmatch(r"\d{4}-\d{2}-\d{2}", rd["date"]): errors.append(("Read Instances", rd["key"], "DateRead", rd["date"]))
for a in authors.values():
    if a["gender"] not in (None, *ALLOWED_GENDERS): errors.append(("Authors", a["id"], "Gender", a["gender"]))
if errors:
    raise RuntimeError("Value check failed: %r" % errors[:10])
print("All value checks passed.")


def style_header(ws, widths):
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="305496")
        c.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def write_sheet(wb, title, headers, rows, sheet_key, key_of, widths):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for row in rows:
        ws.append([row.get(h) for h in headers[:-1]] + [" | ".join(row_notes.get((sheet_key, key_of(row)), [])) or None])
        for ci, h in enumerate(headers, 1):
            p = prov.get((sheet_key, key_of(row), h))
            if p:
                ws.cell(row=ws.max_row, column=ci).fill = KIND_FILL[p[0]]
    style_header(ws, widths)
    return ws


wb = Workbook()
wb.remove(wb.active)
readme = wb.create_sheet("Read Me")

# Authors
author_rows = [{"ImportAuthorID": a["id"], "Author": a["name"], "Country": a["country"], "CountryDetails": a["country_details"],
                "Gender": a["gender"], "BIPOC": a["bipoc"], "LGBTQ+": a["lgbtq"], "Disabled/Neurodiverse": a["disabled"]}
               for a in authors.values()]
write_sheet(wb, "Authors", ["ImportAuthorID", "Author", "Country", "CountryDetails", "Gender", "BIPOC", "LGBTQ+",
                            "Disabled/Neurodiverse", "Review Notes"], author_rows, "Authors",
            lambda r: r["ImportAuthorID"], [14, 32, 22, 38, 10, 8, 8, 12, 70])

# Books (original columns minus Review, then the new ones)
book_rows = [{"ImportBookID": b["id"], "BookTitle": b["title"], "Author": b["author_name"], "ImportAuthorID": b["author_id"],
              "Pages": b["pages"], "Format": b["format"], "Fiction/Nonfiction": b["fnf"], "Genre": b["genre"],
              "OriginalLanguage": b["orig_lang"], "ISBN": b["isbn"], "PublicationYear": b["year"], "Country": b["country_col"],
              "Status": b["status"], "ISBN13": b["isbn13"], "CoverFile": b.get("cover")} for b in books]
write_sheet(wb, "Books", ["ImportBookID", "BookTitle", "Author", "ImportAuthorID", "Pages", "Format", "Fiction/Nonfiction",
                          "Genre", "OriginalLanguage", "ISBN", "PublicationYear", "Country", "Status", "ISBN13", "CoverFile",
                          "Review Notes"], book_rows, "Books", lambda r: r["ImportBookID"],
            [12, 46, 26, 14, 8, 11, 16, 14, 18, 16, 14, 10, 10, 16, 22, 70])

# Read instances (+ Review)
read_rows = [{"Book": r["book_id"], "BookTitle": r["title"], "DateRead": r["date"], "Rating": r["rating"], "Format": r["format"],
              "LanguageRead": r["lang"], "Year": r["year"], "Read": r["read_flag"], "Review": r["review"], "_key": r["key"]} for r in reads]
write_sheet(wb, "Read Instances", ["Book", "BookTitle", "DateRead", "Rating", "Format", "LanguageRead", "Year", "Read", "Review",
                                   "Review Notes"], read_rows, "Read Instances", lambda r: r["_key"],
            [8, 46, 12, 8, 11, 14, 8, 8, 60, 60])

# Problems
ws = wb.create_sheet("Problems")
ph = ["Sheet", "Row ID", "Name / Title", "Issue", "Details"]
ws.append(ph)
for p in problems:
    ws.append([p[h] for h in ph])
style_header(ws, [16, 10, 40, 60, 90])

# Read Me
cover_mb = sum(os.path.getsize(p) for p in cover_files.values()) / 1e6 if cover_files else 0
lines = [
    ("Book database import review", "title"),
    ("Created %s. Your original file was not modified." % RUN_STARTED.strftime("%Y-%m-%d %H:%M"), None),
    ("", None),
    ("How to review", "h"),
    ("Cells your data already had are untouched (apart from the tidy-ups listed below). New values are coloured:", None),
    ("Yellow = looked up on the web.", "web"), ("Orange = a guess or inference - please check.", "guess"),
    ("Blue = moved or derived from your own data (reviews, status).", "derived"),
    ("The 'Review Notes' column on each sheet says where each new value came from. The 'Problems' sheet lists things that need a decision.", None),
    ("Genre, BIPOC, LGBTQ+ and Disabled/Neurodiverse were left exactly as they were.", None),
    ("", None),
    ("What was filled", "h"),
]
for (sheet, col, kind), n in sorted(fill_counts.items()):
    lines.append(("%s - %s: %d %s" % (sheet, col, n, KIND_LABEL[kind]), None))
lines += [("", None), ("Tidy-ups to your own values", "h")]
lines += [("%s: %d" % (k, v), None) for k, v in sorted(changes.items())] or [("(none)", None)]
lines += [("", None), ("Notes", "h"),
          ("Publication year rule used for blanks: %s." % year_rule, None),
          ("Reviews were moved from the Books sheet to the Read Instances sheet (latest read if a book has several).", None),
          ("Books with no Read Instance were given Status 'To Read' (orange) - change any that are DNF or currently reading.", None),
          ("Format values: anything with 'audio' is Audiobook; Kindle/ebook/digital is eBook; everything physical (paperback, hardcover, unknown binding...) is Print.", None),
          ("The empty 'Country' column on the Books sheet was kept as-is.", None),
          ("Covers: %d downloaded, %.1f MB in total (folder: %s)." % (len(cover_files), cover_mb, os.path.basename(os.path.normpath(COVER_DIR))) if DOWNLOAD_COVERS else "Covers: not downloaded.", None),
          ("Problems found: %d (see the Problems sheet)." % len(problems), None),
          ("Network requests this run: %d, answered from cache: %d." % (net.stats["requests"], net.stats["cache_hits"]), None)]
if net.stats["failed"]:
    lines.append(("%d web requests failed this run (timeouts or server errors). Run the notebook again: finished work is cached "
                  "and only the failed requests are retried." % net.stats["failed"], "guess"))
for i, (text, style) in enumerate(lines, 1):
    c = readme.cell(row=i, column=1, value=text)
    if style == "title":
        c.font = Font(bold=True, size=16)
    elif style == "h":
        c.font = Font(bold=True, size=12)
    elif style in KIND_FILL:
        c.fill = KIND_FILL[style]
readme.column_dimensions["A"].width = 130

wb.save(OUTPUT_XLSX)
print("Saved", os.path.abspath(OUTPUT_XLSX))
print("Colour counts:", {KIND_LABEL[k]: sum(n for (s, c, kk), n in fill_counts.items() if kk == k) for k in KIND_FILL})
