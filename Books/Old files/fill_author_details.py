# Fill missing author Country and Gender from Wikipedia - plain-Python version of fill_author_details.ipynb
# Install once:  pip install requests openpyxl
# The '# %%' lines mark cells: VS Code, Spyder and PyCharm run them as notebook cells,
# or just run the whole file with:  python fill_author_details.py
# Edit the settings block near the top first (CONTACT_EMAIL is required).

# %% [markdown]
# # Fill in missing author Country and Gender
#
# Looks up the authors who have a blank **Country** or **Gender** on the Authors sheet of your review workbook,
# using Wikipedia (and the Wikidata facts behind each Wikipedia page).
#
# **Before you start**
# 1. Save this notebook in the same folder as `wordpress_book_database_import_review.xlsx`.
# 2. Put your email address in `CONTACT_EMAIL` below. Wikipedia asks for it so they can contact you if a script misbehaves.
# 3. Run the cells from the top (**Run All**). It takes a few minutes. Run it again any time: answers are remembered.
#
# **What it does and does not do**
# - It only fills **blank** Country and Gender cells. It never overwrites anything you or the earlier notebook entered.
# - It does not touch BIPOC, LGBTQ+, Disabled/Neurodiverse or Country Details.
# - Your original file is never changed. A new file is saved (see `OUTPUT_XLSX`).
# - Hidden columns, hidden/filtered rows and filters are kept exactly as they are. Columns are found by their headings,
#   so hiding them makes no difference.
# - Colours: **yellow** = looked up from a Wikipedia page whose title is exactly the author's name;
#   **orange** = a guess, please check (the name only roughly matched, or the answer was read from the
#   article's wording, or the birth country and citizenship differ).
# - A new sheet, **Author Lookup Log**, lists every answer with a link to the Wikipedia page it came from.

# %%
# ===================== SETTINGS =====================
INPUT_XLSX    = "wordpress_book_database_import_review.xlsx"
OUTPUT_XLSX   = "wordpress_book_database_import_review_authors_filled.xlsx"
CONTACT_EMAIL = "you@example.com"      # <- change this to your email address

FILL_COUNTRY  = True
FILL_GENDER   = True
TEST_LIMIT    = None                   # e.g. 25 to try just a few authors first; None = everyone
CACHE_DB      = "author_lookup_cache.sqlite"

# %%
import collections
import difflib
import json
import os
import re
import sqlite3
import time
import unicodedata

import requests
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

if "@" not in CONTACT_EMAIL or CONTACT_EMAIL == "you@example.com":
    raise RuntimeError("Please put your email address in CONTACT_EMAIL in the settings cell above, then run again.")
if os.path.abspath(INPUT_XLSX) == os.path.abspath(OUTPUT_XLSX):
    raise RuntimeError("OUTPUT_XLSX must be a different file from INPUT_XLSX, so your original is never overwritten.")
if not os.path.exists(INPUT_XLSX):
    raise FileNotFoundError("Could not find %r. Save this notebook in the same folder as the workbook, or change INPUT_XLSX." % INPUT_XLSX)

FILL_YELLOW = PatternFill("solid", fgColor="FFF2CC")
FILL_ORANGE = PatternFill("solid", fgColor="F8CBAD")

COUNTRY_ALIASES = {
    "united states of america": "US", "united states": "US", "usa": "US", "u.s.": "US", "america": "US",
    "united kingdom": "UK", "great britain": "UK", "u.k.": "UK",
    "czechia": "Czech Republic", "state of palestine": "Palestine", "people's republic of china": "China",
    "republic of china": "Taiwan", "united arab emirates": "UAE", "ivory coast": "Côte d'Ivoire",
    "türkiye": "Turkey", "east timor": "Timor-Leste",
}

# Nationality word -> country name (as used in your sheet). Longest phrases are matched first.
DEMONYMS = {
    "American": "US", "Canadian": "Canada", "Québécois": "Canada", "Quebecois": "Canada",
    "British": "UK", "English": "UK", "Scottish": "UK", "Welsh": "UK", "Northern Irish": "UK",
    "Irish": "Ireland", "Australian": "Australia", "New Zealand": "New Zealand", "New Zealander": "New Zealand",
    "French": "France", "German": "Germany", "Italian": "Italy", "Spanish": "Spain", "Portuguese": "Portugal",
    "Dutch": "Netherlands", "Belgian": "Belgium", "Swiss": "Switzerland", "Austrian": "Austria",
    "Swedish": "Sweden", "Norwegian": "Norway", "Danish": "Denmark", "Finnish": "Finland", "Icelandic": "Iceland",
    "Polish": "Poland", "Czech": "Czech Republic", "Slovak": "Slovakia", "Hungarian": "Hungary",
    "Romanian": "Romania", "Bulgarian": "Bulgaria", "Greek": "Greece", "Turkish": "Turkey", "Russian": "Russia",
    "Ukrainian": "Ukraine", "Serbian": "Serbia", "Croatian": "Croatia", "Bosnian": "Bosnia and Herzegovina",
    "Israeli": "Israel", "Palestinian": "Palestine", "Lebanese": "Lebanon", "Egyptian": "Egypt", "Iranian": "Iran",
    "Iraqi": "Iraq", "Syrian": "Syria", "Saudi": "Saudi Arabia", "Afghan": "Afghanistan",
    "Indian": "India", "Pakistani": "Pakistan", "Bangladeshi": "Bangladesh", "Sri Lankan": "Sri Lanka",
    "Nepali": "Nepal", "Chinese": "China", "Japanese": "Japan", "South Korean": "South Korea", "Korean": "South Korea",
    "Vietnamese": "Vietnam", "Thai": "Thailand", "Filipino": "Philippines", "Indonesian": "Indonesia",
    "Malaysian": "Malaysia", "Singaporean": "Singapore", "Taiwanese": "Taiwan",
    "Nigerian": "Nigeria", "Ghanaian": "Ghana", "Kenyan": "Kenya", "South African": "South Africa",
    "Ethiopian": "Ethiopia", "Ugandan": "Uganda", "Zimbabwean": "Zimbabwe", "Somali": "Somalia",
    "Algerian": "Algeria", "Moroccan": "Morocco", "Tunisian": "Tunisia", "Senegalese": "Senegal",
    "Mexican": "Mexico", "Brazilian": "Brazil", "Argentine": "Argentina", "Argentinian": "Argentina",
    "Chilean": "Chile", "Colombian": "Colombia", "Peruvian": "Peru", "Cuban": "Cuba", "Venezuelan": "Venezuela",
    "Haitian": "Haiti", "Jamaican": "Jamaica", "Trinidadian": "Trinidad and Tobago", "Guatemalan": "Guatemala",
}
DEMONYM_RE = re.compile(r"\b(" + "|".join(re.escape(k) for k in sorted(DEMONYMS, key=len, reverse=True)) + r")\b")

WRITER_WORDS = re.compile(r"writer|author|novelist|poet|playwright|dramatist|journalist|essayist|historian|biographer|"
                          r"critic|screenwriter|memoirist|columnist|translator|philosopher|scholar|academic|professor|"
                          r"scientist|theologian|illustrator|cartoonist|lyricist|activist|blogger|comics|"
                          r"psychologist|psychiatrist|economist|sociologist|anthropologist|physician|lawyer|politician", re.I)


# %%
# ===================== helpers =====================
def blank(v):
    return v is None or str(v).strip() == ""


def progress(i, n, label):
    step = max(1, -(-n // 10))
    if i == n or i % step == 0:
        print("  %s: %d of %d" % (label, i, n))


class Net:
    """Polite web requests: remembered answers, a pause between requests, retries. A failed request is not remembered."""
    GAP = {"en.wikipedia.org": 0.15, "query.wikidata.org": 1.0}

    def __init__(self, cache_db, email):
        self.db = sqlite3.connect(cache_db)
        self.db.execute("CREATE TABLE IF NOT EXISTS cache (key TEXT PRIMARY KEY, value TEXT)")
        self.db.commit()
        self.session = requests.Session()
        self.session.headers["User-Agent"] = "BookDatabaseAuthorLookup/1.0 (%s)" % email
        self.last = {}
        self.stats = collections.Counter()

    def get_json(self, url, params=None, headers=None):
        key = "GET %s %s" % (url, json.dumps(params or {}, sort_keys=True))
        row = self.db.execute("SELECT value FROM cache WHERE key=?", (key,)).fetchone()
        if row:
            self.stats["remembered"] += 1
            return json.loads(row[0])
        host = re.sub(r"^https?://", "", url).split("/")[0]
        for attempt in range(4):
            wait = self.last.get(host, 0) + self.GAP.get(host, 0.3) - time.time()
            if wait > 0:
                time.sleep(wait)
            self.last[host] = time.time()
            try:
                r = self.session.get(url, params=params, headers=headers, timeout=40)
            except (requests.ConnectionError, requests.Timeout):
                time.sleep(2 ** attempt)
                continue
            self.stats["requests"] += 1
            if r.status_code == 429 or r.status_code >= 500:
                try:
                    pause = float(r.headers.get("Retry-After", 0))
                except ValueError:
                    pause = 0
                time.sleep(max(pause, 2 ** attempt * 2))
                continue
            try:
                data = r.json() if r.status_code == 200 else {"__status": r.status_code}
            except ValueError:
                data = {"__status": r.status_code}
            self.db.execute("INSERT OR REPLACE INTO cache VALUES (?, ?)", (key, json.dumps(data)))
            self.db.commit()
            return data
        self.stats["failed"] += 1
        return {"__status": -1}          # not remembered, so running again retries it


net = Net(CACHE_DB, CONTACT_EMAIL)


def norm(s):
    s = unicodedata.normalize("NFKD", str(s or ""))
    s = "".join(c for c in s if not unicodedata.combining(c)).lower()
    return re.sub(r"[^a-z0-9]+", "", s)          # "J. R. R. Tolkien" and "J.R.R. Tolkien" both become "jrrtolkien"


def title_matches(name, title):
    """'exact', 'fuzzy' or None - how well a Wikipedia page title matches the author's name."""
    t = re.sub(r"\s*\([^)]*\)\s*$", "", title)
    a, b = norm(name), norm(t)
    if not a or not b:
        return None
    if a == b:
        return "exact"
    return "fuzzy" if difflib.SequenceMatcher(None, a, b).ratio() >= 0.93 else None


# %%
# ===================== Wikipedia and Wikidata lookups =====================
WIKI_API = "https://en.wikipedia.org/w/api.php"
SPARQL = "https://query.wikidata.org/sparql"


def wikipedia_search(name):
    data = net.get_json(WIKI_API, {
        "action": "query", "format": "json", "formatversion": 2, "generator": "search",
        "gsrsearch": '"%s"' % name, "gsrnamespace": 0, "gsrlimit": 5,
        "prop": "extracts|pageprops", "exintro": 1, "explaintext": 1, "exlimit": 5, "ppprop": "wikibase_item",
    })
    pages = (data.get("query") or {}).get("pages") or [] if isinstance(data, dict) else []
    return sorted(pages, key=lambda p: p.get("index", 99))


def pick_page(name, pages):
    """Returns (page, 'exact'/'fuzzy') for the one page that is clearly about this writer, else (None, reason)."""
    good = []
    for p in pages:
        extract = p.get("extract") or ""
        how = title_matches(name, p.get("title", ""))
        if not how or "may refer to" in extract[:200] or not WRITER_WORDS.search(extract[:500]):
            continue
        good.append((p, how))
    exact = [g for g in good if g[1] == "exact"]
    if len(exact) == 1:
        return exact[0]
    if len(exact) > 1:
        return None, "several Wikipedia pages share this name"
    if len(good) == 1:
        return good[0]
    return None, "no matching Wikipedia page" if not good else "several possible Wikipedia pages"


def sparql(query):
    data = net.get_json(SPARQL, {"query": query, "format": "json"}, headers={"Accept": "application/sparql-results+json"})
    return data.get("results", {}).get("bindings", []) if isinstance(data, dict) else []


def wikidata_facts(qids):
    facts = {}
    qids = sorted(set(qids))
    batches = [qids[i:i + 80] for i in range(0, len(qids), 80)]
    for n, batch in enumerate(batches, 1):
        progress(n, len(batches), "Wikidata batches")
        q = ("SELECT ?a ?genderLabel ?citLabel ?bcLabel WHERE { VALUES ?a { %s } "
             "OPTIONAL { ?a wdt:P21 ?gender . } OPTIONAL { ?a wdt:P27 ?cit . } "
             "OPTIONAL { ?a wdt:P19 ?bp . ?bp wdt:P17 ?bc . } "
             'SERVICE wikibase:label { bd:serviceParam wikibase:language "en" . } }') % " ".join("wd:" + x for x in batch)
        for row in sparql(q):
            qid = row["a"]["value"].rsplit("/", 1)[-1]
            f = facts.setdefault(qid, {"gender": set(), "citizenship": set(), "birth_country": set()})
            for src, dst in (("genderLabel", "gender"), ("citLabel", "citizenship"), ("bcLabel", "birth_country")):
                v = row.get(src, {}).get("value", "")
                if v and not v.startswith("http") and not re.fullmatch(r"Q\d+", v):
                    f[dst].add(v)
    return facts


def decide_gender(f):
    mapped = set()
    for g in f["gender"]:
        g = g.lower()
        mapped.add("Male" if g in ("male", "trans man") else "Female" if g in ("female", "trans woman") else "Other")
    return mapped.pop() if len(mapped) == 1 else None


def decide_country(f):
    """(country or None, note, born-elsewhere?)"""
    cit = sorted({COUNTRY_ALIASES.get(c.lower(), c) for c in f["citizenship"]})
    born = sorted({COUNTRY_ALIASES.get(c.lower(), c) for c in f["birth_country"]})
    if len(cit) == 1:
        diff = bool(born) and born != cit
        return cit[0], "citizenship %s%s" % (cit[0], "; born in " + ", ".join(born) if diff else ""), diff
    if len(born) == 1:
        return born[0], "born in %s (citizenship: %s)" % (born[0], ", ".join(cit) or "unknown"), len(cit) > 1
    return None, "", False


# ---- reading the article's own wording, for authors Wikidata has no answer for ----
def gender_from_text(text):
    t = (text or "").lower()
    if re.search(r"non-?binary|genderqueer|genderfluid|\bagender\b|they/them", t):
        return "Other"
    he = len(re.findall(r"\b(he|his|him|himself)\b", t))
    she = len(re.findall(r"\b(she|her|hers|herself)\b", t))
    if he >= 2 and he >= 3 * she:
        return "Male"
    if she >= 2 and she >= 3 * he:
        return "Female"
    return None


def country_from_text(text):
    """The nationality word(s) in front of the first job title, e.g. 'is a Nigerian-American novelist' -> US."""
    text = text or ""
    job = WRITER_WORDS.search(text)
    if not job:
        return None
    found = DEMONYM_RE.findall(text[:job.start()][-160:])
    return DEMONYMS[found[-1]] if found else None


# %%
# ===================== read the workbook =====================
wb = load_workbook(INPUT_XLSX)              # keeps formatting, hidden columns/rows and filters
if "Authors" not in wb.sheetnames:
    raise RuntimeError("The workbook has no sheet called 'Authors'.")
ws = wb["Authors"]

cols = {}
for c in ws[1]:
    if not blank(c.value):
        cols[str(c.value).strip()] = c.column
for needed in ("Author", "Country", "Gender"):
    if needed not in cols:
        raise RuntimeError("The Authors sheet has no '%s' column." % needed)
if "Review Notes" not in cols:
    cols["Review Notes"] = ws.max_column + 1
    ws.cell(1, cols["Review Notes"], "Review Notes").font = Font(bold=True)

hidden = [k for k, d in ws.column_dimensions.items() if d.hidden]
print("Authors sheet: %d rows. Hidden columns (kept as they are): %s" % (ws.max_row - 1, ", ".join(hidden) or "none"))

todo = []
for r in range(2, ws.max_row + 1):
    name = ws.cell(r, cols["Author"]).value
    if blank(name):
        continue
    want_c = FILL_COUNTRY and blank(ws.cell(r, cols["Country"]).value)
    want_g = FILL_GENDER and blank(ws.cell(r, cols["Gender"]).value)
    if want_c or want_g:
        todo.append({"row": r, "name": str(name).strip(), "want_c": want_c, "want_g": want_g,
                     "id": ws.cell(r, cols["ImportAuthorID"]).value if "ImportAuthorID" in cols else None})

skipped_multi = [t for t in todo if re.search(r"\s&\s|\sand\s|,|\bvarious\b|\banonymous\b|\bet al\b", t["name"], re.I)]
todo = [t for t in todo if t not in skipped_multi]
if TEST_LIMIT:
    todo = todo[:TEST_LIMIT]
print("%d authors have a blank Country or Gender to look up (%d skipped: several people or no real name)."
      % (len(todo), len(skipped_multi)))

# %%
# ===================== search Wikipedia =====================
for n, t in enumerate(todo, 1):
    progress(n, len(todo), "Wikipedia searches")
    page, how = pick_page(t["name"], wikipedia_search(t["name"]))
    t["page"], t["how"] = (page, how) if page else (None, how)
found = [t for t in todo if t["page"]]
print("Matched a Wikipedia page for %d of %d authors." % (len(found), len(todo)))

facts = wikidata_facts([t["page"]["pageprops"]["wikibase_item"] for t in found
                        if (t["page"].get("pageprops") or {}).get("wikibase_item")])
print("Wikidata facts for %d people." % len(facts))

# %%
# ===================== decide the answers =====================
results = []     # one row per filled-in value, for the log sheet
for t in todo:
    t["gender"] = t["country"] = None
    page = t["page"]
    if not page:
        continue
    qid = (page.get("pageprops") or {}).get("wikibase_item")
    f = facts.get(qid) if qid else None
    exact = t["how"] == "exact"
    url = "https://en.wikipedia.org/wiki/" + page["title"].replace(" ", "_")
    intro = page.get("extract") or ""

    if t["want_g"]:
        g, source, sure = (decide_gender(f) if f else None), None, exact
        if g:
            source = "Wikidata %s" % qid
        else:
            g = gender_from_text(intro)
            if g:
                source, sure = "pronouns in the Wikipedia introduction", False
        if g:
            t["gender"] = (g, source, sure)
    if t["want_c"]:
        c, source, sure = None, None, exact
        if f:
            c, note, born_elsewhere = decide_country(f)
            if c:
                source, sure = "Wikidata %s: %s" % (qid, note), exact and not born_elsewhere
        if not c:
            c = country_from_text(intro)
            if c:
                source, sure = "nationality word in the Wikipedia introduction", False
        if c:
            t["country"] = (c, source, sure)

# %%
# ===================== write the new workbook =====================
filled_c = filled_g = 0
log_rows = []
for t in todo:
    r, page = t["row"], t["page"]
    url = ("https://en.wikipedia.org/wiki/" + page["title"].replace(" ", "_")) if page else ""
    notes = []
    for field, key in (("Gender", "gender"), ("Country", "country")):
        if t[key]:
            value, source, sure = t[key]
            cell = ws.cell(r, cols[field])
            if blank(cell.value):                       # re-check: never overwrite
                cell.value = value
                cell.fill = FILL_YELLOW if sure else FILL_ORANGE
                notes.append("%s: %s (Wikipedia: %s%s)" % (field, source, page["title"], "" if t["how"] == "exact" else ", name only roughly matched"))
                log_rows.append((t["id"], t["name"], field, value, source, page["title"], url, "Looked up" if sure else "Guess - check"))
                if field == "Country":
                    filled_c += 1
                else:
                    filled_g += 1
    if notes:
        nc = ws.cell(r, cols["Review Notes"])
        nc.value = ((str(nc.value) + " | ") if not blank(nc.value) else "") + " | ".join(notes)

if "Author Lookup Log" in wb.sheetnames:
    del wb["Author Lookup Log"]
log = wb.create_sheet("Author Lookup Log")
head = ["ImportAuthorID", "Author", "Field", "Value", "Based on", "Wikipedia page", "Link", "Confidence"]
log.append(head)
for c in log[1]:
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor="305496")
for row in log_rows:
    log.append(list(row))
    link = log.cell(log.max_row, 7)
    link.hyperlink = row[6]
    link.font = Font(color="0563C1", underline="single")
    log.cell(log.max_row, 8).fill = FILL_YELLOW if row[7] == "Looked up" else FILL_ORANGE
for letter, width in zip("ABCDEFGH", (14, 30, 10, 18, 52, 30, 50, 14)):
    log.column_dimensions[letter].width = width
log.freeze_panes = "A2"
log.auto_filter.ref = log.dimensions

# authors that could not be answered, with the reason
nf = [t for t in todo if not (t["gender"] or t["country"])]
log.append([])
log.append(["Not found", "", "", "", "", "", "", ""])
log.cell(log.max_row, 1).font = Font(bold=True)
for t in nf:
    log.append([t["id"], t["name"], "", "", t["how"] if not t["page"] else "Wikipedia page found, but it gave no usable answer", "", "", ""])

wb.save(OUTPUT_XLSX)

print()
print("Saved %s" % OUTPUT_XLSX)
print("Filled %d Gender and %d Country cells (yellow = looked up, orange = please check)." % (filled_g, filled_c))
print("%d authors could not be answered (listed at the bottom of the 'Author Lookup Log' sheet)." % len(nf))
print("Requests: %d new, %d remembered, %d failed." % (net.stats["requests"], net.stats["remembered"], net.stats["failed"]))
if net.stats["failed"]:
    print("Some requests failed (usually a connection hiccup). Run the notebook again to retry just those.")
